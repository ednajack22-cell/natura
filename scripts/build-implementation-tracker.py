from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path
from zipfile import ZipFile

OUT = "Natura_Suivi_Implementation_Priorites.docx"
NAVY = "163D2D"
GREEN = "397550"
PALE = "E8F1EA"
LIGHT = "F3F5F3"
GOLD = "9A6B2F"
RED = "9A4934"
MUTED = "65736B"
INK = "24352C"
WHITE = "FFFFFF"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18
for name, size, color, before, after in [
    ("Heading 1", 17, NAVY, 16, 8), ("Heading 2", 13.5, GREEN, 12, 6), ("Heading 3", 11.5, NAVY, 9, 4)
]:
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

def set_font(run, size=10.5, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc.get_or_add_tcPr()
    tc_mar = tc.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc.append(tc_mar)
    for side, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)

def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")

def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, 9, WHITE, True)
    mark_header_row(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if i == 0 and str(value).startswith("[x]"):
                shade(cells[i], PALE)
            elif i == 0 and str(value).startswith("[~]"):
                shade(cells[i], "FFF2D9")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, 8.7, GREEN if str(value).startswith("[x]") else INK, i == 0)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(hp.add_run("NATURA  /  PILOTAGE PRODUIT"), 8.5, MUTED, True)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(fp.add_run("Document vivant - mise a jour du 10 aout 2026"), 8, MUTED)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run("SUIVI D'IMPLEMENTATION"), 9, GOLD, True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
set_font(p.add_run("Natura - priorites, realisations et prochaine execution"), 25, NAVY, True, name="Aptos Display")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
set_font(p.add_run("Tableau de bord de reference pour ne perdre aucune fonctionnalite, verifier les livrables et suivre l'avancement jusqu'a la mise en ligne."), 11.5, MUTED)

meta = doc.add_table(rows=2, cols=4)
meta.style = "Table Grid"
mark_header_row(meta.rows[0])
meta_rows = [("Version", "2.0", "Prochain lot", "Mediatheque avancee"), ("Environnement", "Local / Supabase", "Etat global", "Construction - environ 90 %")]
for ri, row in enumerate(meta_rows):
    for ci, value in enumerate(row):
        cell = meta.rows[ri].cells[ci]
        shade(cell, LIGHT if ci % 2 == 0 else WHITE)
        r = cell.paragraphs[0].add_run(value)
        set_font(r, 8.8, MUTED if ci % 2 == 0 else INK, ci % 2 == 0)
set_table_geometry(meta, [1500, 2100, 1500, 4260])

doc.add_heading("1. Regles de suivi", level=1)
for text in [
    "[x] REALISE : implemente et valide par compilation ou test adapte.",
    "[~] EN COURS : implementation active dans le lot courant.",
    "[ ] A FAIRE : recommande, mais pas encore livre.",
    "Une fonction n'est marquee realisee que si son parcours visible et son stockage sont verifies.",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), 10, GREEN if text.startswith("[x]") else GOLD if text.startswith("[~]") else INK, text.startswith("["))

doc.add_heading("2. Checklist des parties realisees", level=1)
completed = [
    ("[x]", "Architecture", "Next.js, TypeScript, Supabase local, PostgreSQL et Docker local operationnels.", "Build + 74 tests DB"),
    ("[x]", "Design public", "Accueil premium, responsive, navigation complete et fiches de plantes.", "Controle navigateur"),
    ("[x]", "Administration", "Dashboard prive, navigation groupee, bibliotheque filtrable et creation de contenu.", "Session admin protegee"),
    ("[x]", "Edition", "Edition des brouillons et articles publies, apercu, statuts et historique.", "Revisions PostgreSQL"),
    ("[x]", "Navigation editoriale", "Sections et sous-sections du plan reliees a l'editeur.", "Taxonomies partagees"),
    ("[x]", "Verticale plantes", "Champs structures, pages publiques, accueil, recherche et sitemap depuis PostgreSQL.", "3 fiches migrees"),
    ("[x]", "SEO", "Rapport avant publication, score, mots-cles, intention et derogation explicite.", "Cherry tomato 84/100"),
    ("[x]", "Mediatheque", "Import, validation, alt, credit, licence, dimensions et selection dans l'editeur.", "Actif Supabase verifie"),
    ("[x]", "Diagnostics", "Editeur structure, filtres publics, fiche complete, relations avec les plantes et SEO.", "Parcours admin-public verifie"),
    ("[x]", "Guides editoriaux", "Editeur structure, apercu, bibliotheque filtree, HowTo, recherche et relations plantes.", "Parcours admin-public verifie"),
    ("[x]", "Calendriers dynamiques", "Editeur douze mois, vue annuelle, recherche, relations plantes et PDF annuel ou personnalise.", "Exports 1, 2 et 12 mois verifies"),
    ("[x]", "Comparatifs et avis", "Modele, scorecard, editeur, apercu, pages publiques, recherche, sitemap et transparence commerciale.", "51 tests DB + parcours UI"),
    ("[x]", "Pages institutionnelles", "A propos, politique editoriale, auteurs, contact, mentions legales, confidentialite et transparence commerciale.", "7 pages + 56 tests DB"),
    ("[x]", "Workflow editorial", "Auteurs, reviseurs, approbation, programmation, calendrier reel, comparaison et restauration sure des versions.", "66 tests DB + parcours UI"),
    ("[x]", "Recherche PostgreSQL", "FTS pondere, pertinence, prefixes, suggestions, filtres et liens vers toutes les verticales publiees.", "74 tests DB + parcours UI"),
    ("[x]", "Securite", "RLS, controle du role admin, en-tetes HTTP et routes admin non indexables.", "Politiques DB testees"),
    ("[x]", "Qualite", "ESLint, TypeScript, compilation de production et tests PostgreSQL.", "Dernier controle reussi"),
]
add_table(["Statut", "Domaine", "Livraison", "Preuve"], completed, [850, 1450, 4700, 2360])

doc.add_heading("3. Registre des priorites restantes", level=1)
priorities = [
    ("[x] P0", "Problemes et diagnostics", "Schema, editeur specialise, pages publiques, symptomes, causes, actions et liens plantes.", "Livre en version 1.1"),
    ("[x] P0", "Verticales prioritaires", "Diagnostics, guides, calendriers, comparatifs et pages institutionnelles sont livres et relies au dashboard.", "Lot structurel termine"),
    ("[ ] P0", "Validation institutionnelle", "Confirmer operateur, juridiction, adresse, email public, traitements reels et validation juridique avant domaine public.", "Avant mise en ligne"),
    ("[x] P1", "Workflow avance", "Auteurs, reviseurs, approbation, programmation, calendrier reel, comparaison et restauration en brouillon.", "Livre en version 1.9"),
    ("[ ] P1", "Mediatheque avancee", "Modifier, remplacer, supprimer avec controle d'usage, recadrer et gerer les brouillons prives.", "A planifier"),
    ("[x] P1", "Recherche PostgreSQL", "FTS tous contenus, pertinence, suggestions, filtres et liens internes.", "Livre en version 2.0"),
    ("[ ] P1", "Outils interactifs", "Plant finder, pot, arrosage, calendrier, gel, planner et diagnostic guide.", "Backlog"),
    ("[ ] P1", "SEO avance", "Breadcrumbs, schemas par type, cannibalisation, liens casses, auteur et fraicheur.", "Backlog"),
    ("[ ] P1", "Contenu lancement", "15 plantes, 10-15 problemes, 5-10 guides, calendriers, outils, sources et pages legales.", "Editorial"),
    ("[ ] P2", "Qualite production", "Tests E2E, accessibilite, performance, CSP, sauvegardes et monitoring.", "Avant hebergement"),
    ("[ ] P2", "Mise en ligne", "Hebergement, domaine, HTTPS, CI/CD, Search Console, GA4 et consentement.", "Derniere phase"),
    ("[ ] DEC", "Langue cible", "Confirmer anglais seul ou architecture multilingue anglais/francais.", "Decision requise"),
]
add_table(["Priorite", "Chantier", "Perimetre", "Etape"], priorities, [1000, 1800, 4840, 1720])

doc.add_heading("4. Suivi des verticales editoriales", level=1)
doc.add_heading("4.1 Problemes et diagnostics - realise", level=2)
doc.add_paragraph("Objectif : permettre de creer un probleme de plante structure dans l'administration, le publier et l'afficher dans une bibliotheque publique reliee aux fiches de plantes.")
diagnostic_steps = [
    ("[x] D1", "Modele PostgreSQL", "Symptomes, causes, verification, actions, urgence, plantes concernees et sources.", "Migration + RLS"),
    ("[x] D2", "Editeur specialise", "Champs structures et controle de completude avant publication.", "Interface admin"),
    ("[x] D3", "Bibliotheque publique", "Liste filtrable des problemes avec cartes et niveaux d'urgence.", "/plant-care/problems"),
    ("[x] D4", "Fiche diagnostic", "Symptomes, causes probables, verification, plan d'action et avertissements.", "/plant-care/problems/[slug]"),
    ("[x] D5", "Relations plantes", "Afficher les problemes lies depuis les fiches de plantes et inversement.", "Maillage interne"),
    ("[x] D6", "SEO et schemas", "Metadata, canonical, Article et mots-cles dans le rapport SEO.", "Rapport SEO"),
    ("[x] D7", "Validation", "Tests DB, TypeScript, build, mobile et parcours admin-public.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], diagnostic_steps, [950, 1800, 4710, 1900])

doc.add_heading("4.2 Guides editoriaux - realise", level=2)
doc.add_paragraph("Objectif : creer des projets pratiques structures dans l'administration, les previsualiser, les publier et les relier a la navigation, aux plantes et a la recherche globale.")
guide_steps = [
    ("[x] G1", "Modele PostgreSQL", "Type, difficulte, duree, public, prerequis, materiel, etapes, conseils et limites.", "Migration + RLS"),
    ("[x] G2", "Editeur specialise", "Champs structures, image, plantes liees et controle avant publication.", "Interface admin"),
    ("[x] G3", "Apercu prive", "Materiel, preparation et methode visibles avant publication.", "Apercu admin"),
    ("[x] G4", "Bibliotheque publique", "Liste filtrable par type avec duree, difficulte et nombre d'etapes.", "/small-space-gardening/guides"),
    ("[x] G5", "Fiche guide", "HowTo, prerequis, materiel, etapes, conseils, limites et plantes associees.", "/guides/[slug]"),
    ("[x] G6", "Decouverte", "Menu, page Small-Space, recherche globale, sitemap et maillage interne.", "Navigation + SEO"),
    ("[x] G7", "Validation", "33 tests DB, TypeScript, ESLint, build et parcours admin-public.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], guide_steps, [950, 1800, 4710, 1900])

doc.add_heading("4.3 Calendriers dynamiques - realise", level=2)
doc.add_paragraph("Objectif : administrer une planification complete sur douze mois, contextualisee par climat, hemisphere et espace de culture, puis la relier aux plantes et aux outils de decouverte du site.")
calendar_steps = [
    ("[x] C1", "Modele PostgreSQL", "Profil de contexte, activites mensuelles ordonnees, avertissements et relations plantes.", "Migration + RLS"),
    ("[x] C2", "Editeur specialise", "Contexte, image, douze panneaux mensuels, priorites et plantes associees.", "Interface admin"),
    ("[x] C3", "Apercu prive", "Contexte et couverture des douze mois visibles avant publication.", "Apercu admin"),
    ("[x] C4", "Bibliotheque publique", "Filtres par climat et espace avec indicateurs de couverture et de taches.", "/planting-calendar/calendars"),
    ("[x] C5", "Calendrier consultable", "Vue annuelle, selection des douze mois, activites, priorites, avertissements et plantes liees.", "/calendars/[slug]"),
    ("[x] C6", "PDF et decouverte", "Export annuel ou selection libre des mois, menu, recherche, sitemap, schemas et maillage plantes.", "PDF 1, 2 ou 12 mois"),
    ("[x] C7", "Validation", "41 tests DB, TypeScript, ESLint, build et parcours admin-public.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], calendar_steps, [950, 1800, 4710, 1900])

doc.add_heading("4.4 Comparatifs et avis - realise", level=2)
doc.add_paragraph("Objectif : publier des comparatifs utiles et transparents, avec niveau de preuve, methode, relation commerciale, criteres ponderes, scores, forces, limites et verdict explicites.")
review_steps = [
    ("[x] R1", "Modele PostgreSQL", "Profils, criteres ponderes, options, scores, divulgation commerciale et niveau de preuve.", "Migration + RLS pretes"),
    ("[x] R2", "Editeur specialise", "Criteres dynamiques, total 100 %, options, forces, limites, scores et recommandation.", "Interface admin"),
    ("[x] R3", "Controle publication", "Methode, divulgation, verdict, trois criteres, deux options et matrice complete obligatoires.", "Blocage ou derogation"),
    ("[x] R4", "Apercu prive", "Methode, divulgation, scorecard et verdict visibles avant publication.", "Apercu admin"),
    ("[x] R5", "Bibliotheque et fiche", "Filtres par categorie, tableau comparatif, details, methode et declaration commerciale visible.", "/reviews + /reviews/[slug]"),
    ("[x] R6", "SEO et decouverte", "Canonical, robots, breadcrumbs, Article, ItemList, menu, recherche et sitemap.", "Integration complete"),
    ("[x] R7", "Activation et validation DB", "Migration appliquee, exemple source, types regeneres, 51 tests PostgreSQL et parcours desktop/mobile valides.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], review_steps, [950, 1800, 4710, 1900])

doc.add_heading("4.5 Pages institutionnelles - realise", level=2)
doc.add_paragraph("Objectif : administrer et publier un centre de confiance coherent, avec pages editoriales, legales et commerciales structurees, tout en rendant visibles les validations encore requises avant la mise en ligne.")
institutional_steps = [
    ("[x] I1", "Modele PostgreSQL", "Type de page, sections, contact, date de revision, notice et indicateur de validation juridique.", "Migration + RLS"),
    ("[x] I2", "Editeur specialise", "Sections dynamiques, ordre, contact, date de revision et controle avant publication.", "Dashboard professionnel"),
    ("[x] I3", "Sept pages publiques", "A propos, politique editoriale, auteurs, contact, mentions legales, confidentialite et divulgation commerciale.", "Routes dynamiques"),
    ("[x] I4", "Transparence pre-lancement", "Avertissements visibles pour identite operateur, juridiction, contact et traitements de donnees a confirmer.", "Risque signale"),
    ("[x] I5", "Decouverte", "Footer, recherche globale, sitemap et rubrique institutionnelle dans l'administration.", "Navigation complete"),
    ("[x] I6", "SEO et schemas", "Metadata, canonical, breadcrumbs et schemas AboutPage, ContactPage ou WebPage.", "Template public"),
    ("[x] I7", "Validation", "56 tests DB, TypeScript, ESLint, build, desktop, mobile, console et sitemap.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], institutional_steps, [950, 1800, 4710, 1900])

doc.add_heading("4.6 Workflow editorial avance - realise", level=2)
doc.add_paragraph("Objectif : organiser la responsabilite editoriale, separer la redaction de la validation, programmer les publications et conserver un historique restaurable sans republier accidentellement un contenu.")
workflow_steps = [
    ("[x] W1", "Contributeurs", "Annuaire professionnel des auteurs et reviseurs, expertises, conflits, disponibilite et profil public.", "Dashboard protege"),
    ("[x] W2", "Affectation", "Auteur responsable et reviseur distinct attribues depuis l'editeur avec notes de verification.", "Workflow PostgreSQL"),
    ("[x] W3", "Approbation", "Soumission, demandes de correction, approbation datee et invalidation automatique si les responsables changent.", "Audit editorial"),
    ("[x] W4", "Programmation", "Date future autorisee uniquement apres approbation et controle de completude avant publication.", "Blocage securise"),
    ("[x] W5", "Calendrier reel", "Navigation mensuelle, file des contenus dus et publication volontaire des seuls contenus approuves.", "/admin/calendar"),
    ("[x] W6", "Versions", "Comparaison champ par champ et restauration protegee qui replace toujours l'article en brouillon prive.", "Historique + RPC"),
    ("[x] W7", "Validation", "66 tests PostgreSQL/RLS, TypeScript, ESLint, build, desktop, mobile et console navigateur.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], workflow_steps, [950, 1800, 4710, 1900])

doc.add_heading("4.7 Recherche PostgreSQL avancee - realise", level=2)
doc.add_paragraph("Objectif : remplacer le filtrage en memoire par une recherche PostgreSQL classee, limitee aux contenus publies et capable de guider le lecteur vers toutes les verticales de Natura.")
search_steps = [
    ("[x] S1", "Document FTS", "Vecteur stocke pondere : titre et SEO en A, resume en B, contenu structure en C.", "Colonne tsvector"),
    ("[x] S2", "Index et requetes", "Index GIN partiel sur les contenus publies, websearch et recherche par prefixe.", "Migration PostgreSQL"),
    ("[x] S3", "Classement", "Pertinence textuelle, bonus titre exact ou partiel, puis fraicheur et titre.", "RPC classee"),
    ("[x] S4", "Suggestions", "Requetes editoriales et titres publies proposes dans le formulaire et en liens associes.", "Datalist + liens"),
    ("[x] S5", "Filtres", "Tous contenus, plantes, diagnostics, guides, calendriers, avis, outils et centre de confiance.", "/search"),
    ("[x] S6", "Etats et accessibilite", "Requete courte, resultat vide, indisponibilite DB, filtres nommes et navigation clavier.", "Parcours public"),
    ("[x] S7", "Validation", "74 tests PostgreSQL/RLS, TypeScript, ESLint, build, rendu desktop et console navigateur.", "Definition of done"),
]
add_table(["ID", "Livrable", "Critere", "Sortie"], search_steps, [950, 1800, 4710, 1900])

doc.add_heading("5. Definition de termine", level=1)
for item in [
    "Le contenu peut etre cree, modifie, previsualise et publie par un administrateur.",
    "Les donnees publiees viennent de PostgreSQL et respectent les politiques d'acces.",
    "Le contenu apparait dans la bonne rubrique et possede une URL publique stable.",
    "Le rapport SEO et les controles de qualite sont disponibles avant publication.",
    "Le parcours fonctionne sur ordinateur et mobile sans erreur console de l'application.",
    "Les tests adaptes, le typage et la compilation de production reussissent.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(item)

doc.add_heading("6. Journal des mises a jour", level=1)
log_rows = [
    ("10 aout 2026", "2.0", "Recherche PostgreSQL avancee livree : tsvector pondere, index GIN, pertinence, prefixes, suggestions, filtres, resultats inter-verticales et etats accessibles. 74 tests DB, build et parcours navigateur valides."),
    ("10 aout 2026", "1.9", "Workflow editorial avance livre : contributeurs, affectation auteur-reviseur, approbation, programmation, calendrier reel, comparaison et restauration sure. 66 tests DB, build et parcours responsive valides."),
    ("3 aout 2026", "1.8", "Pages institutionnelles livrees : sept pages dynamiques, editeur structure, controles juridiques pre-lancement, footer, recherche, sitemap, SEO et schemas. 56 tests DB et build valides."),
    ("3 aout 2026", "1.7", "Comparatifs actives et verifies : migration locale appliquee, 51 tests DB, sources visibles, SEO 90/100, responsive corrige et aucun message en console navigateur."),
    ("3 aout 2026", "1.6", "Verticale comparatifs codee : schema, RLS, editeur professionnel, controle avant publication, apercu, pages publiques, SEO, recherche et sitemap. Build valide; 51 tests DB prepares, execution en attente de Docker."),
    ("3 aout 2026", "1.5", "Export PDF personnalise : selection d'un ou plusieurs mois, PDF sans pages inutiles, choix total/effacement et export annuel conserve."),
    ("3 aout 2026", "1.4", "Calendrier enrichi : vue annuelle consultable, navigation mensuelle et export PDF direct avec couverture et douze pages mensuelles, rendu verifie."),
    ("3 aout 2026", "1.3", "Verticale calendriers livree : base, editeur douze mois, apercu, filtres, vue mensuelle, SEO, recherche, sitemap et relations plantes. 41 tests DB valides."),
    ("3 aout 2026", "1.2", "Verticale guides livree : base, editeur, apercu, filtres, HowTo, recherche, sitemap et relations plantes. 33 tests DB valides."),
    ("3 aout 2026", "1.1", "Verticale diagnostics livree : base, editeur, filtres, pages publiques, SEO et relations plantes. 27 tests DB valides."),
    ("3 aout 2026", "1.0", "Creation du registre. SEO mots-cles et mediatheque marques realises. Verticale diagnostics ouverte."),
]
add_table(["Date", "Version", "Modification"], log_rows, [1500, 1100, 6760])

doc.core_properties.title = "Natura - Suivi d'implementation des priorites"
doc.core_properties.subject = "Checklist des realisations et feuille de route"
doc.core_properties.author = "Natura Editorial Studio"
doc.save(OUT)
output_path = Path(OUT)
with ZipFile(output_path) as archive:
    assert archive.testzip() is None, "Le fichier DOCX est endommage."
verified = Document(output_path)
document_text = "\n".join(paragraph.text for paragraph in verified.paragraphs)
table_text = "\n".join(cell.text for table in verified.tables for row in table.rows for cell in row.cells)
for required in ["Checklist des parties realisees", "Registre des priorites restantes", "[x] D7", "[x] G7", "[x] C7", "[x] R7", "[x] I7", "[x] W7", "[x] S7", "74 tests DB", "SEO 90/100", "Exports 1, 2 et 12 mois", "Workflow editorial avance", "Recherche PostgreSQL avancee", "Mediatheque avancee", "Validation institutionnelle"]:
    assert required in f"{document_text}\n{table_text}", f"Element absent du document : {required}"
print(f"{OUT} | controle structurel OK | {len(verified.paragraphs)} paragraphes | {len(verified.tables)} tableaux")
