from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

ROOT = Path(r"C:\Users\NEIL\Desktop\__Natura_Web_Site")
source = ROOT / "Natura_Plan_Detaille_Realisation.docx"
output = ROOT / "Natura_Plan_Detaille_Realisation_v1.1.docx"

doc = Document(source)

if doc.tables and len(doc.tables[0].rows) > 1:
    doc.tables[0].cell(1, 2).text = "1.1"
    doc.tables[0].cell(1, 3).text = "2 août 2026"

doc.add_page_break()
doc.add_heading("15. Audit senior, benchmark et réalignement du studio", level=1)
p = doc.add_paragraph()
r = p.add_run("Version 1.1 — état vérifié le 2 août 2026")
r.bold = True
r.font.size = Pt(11)

doc.add_paragraph(
    "Cette section met à jour le plan directeur après l’audit du site, du dashboard, "
    "du parcours de création, de la base locale et des références professionnelles de CMS."
)

doc.add_heading("15.1 Diagnostic de l’erreur d’hydratation", level=2)
doc.add_paragraph(
    "La trace fournie contient les attributs bis_register et bis_skin_checked, injectés par "
    "une extension de navigateur avant l’hydratation React. Next.js répertorie explicitement "
    "les extensions modifiant le HTML comme cause possible. Le code Natura ne contient pas "
    "de branche navigateur, valeur aléatoire ou structure HTML invalide expliquant cette trace."
)
for text in [
    "Tolérer les attributs externes uniquement au niveau des racines HTML.",
    "Ne pas masquer globalement les vraies erreurs de React.",
    "Désactiver pour localhost l’extension qui injecte les attributs bis_* si l’avertissement persiste.",
    "Conserver lint, TypeScript, build de production et tests navigateur comme contrôles de référence.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("15.2 Écarts confirmés par rapport au plan", level=2)
for text in [
    "Les sous-menus Grow, Small-Space Gardening, Plant Care, Planting Calendar, Tools et Reviews n’étaient pas implémentés.",
    "La création d’un brouillon retournait au tableau de bord au lieu d’ouvrir le contenu créé.",
    "La bibliothèque ne permettait pas d’ouvrir un éditeur complet.",
    "Le contenu n’était pas rattaché à la section et à la sous-section du menu public.",
    "Il n’existait ni aperçu privé, ni historique de révisions, ni journal éditorial.",
    "Les fiches publiques restent encore alimentées par des données TypeScript statiques, distinctes de PostgreSQL.",
    "La médiathèque ne permet pas encore l’upload, la licence, le crédit et la sélection depuis l’éditeur.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("15.3 Benchmark professionnel retenu", level=2)
benchmarks = [
    ("Sanity Studio", "Sépare dashboard, structure documentaire et présentation; organise listes, vues, menus et workflows; ajoute validation, aperçu et releases."),
    ("Directus", "Traite les collections comme un module de gestion avec recherche, tri, layouts, pages item et versions non publiées."),
    ("Contentful", "Propose un aperçu des brouillons dans l’éditeur ou dans un nouvel onglet, séparé de la diffusion publique."),
    ("Supabase/PostgreSQL", "Fournit RLS pour les fichiers et données, Storage sécurisé et recherche plein texte indexable."),
]
for name, detail in benchmarks:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(f"{name}. ")
    run.bold = True
    p.add_run(detail)

doc.add_heading("15.4 Correctifs et améliorations intégrés", level=2)
for text in [
    "Méga-menus publics et navigation mobile avec toutes les sous-rubriques prévues.",
    "Rubriques publiques complétées avec ancres stables.",
    "Sidebar admin groupée par workflow éditorial, base de connaissances et opérations.",
    "Filtres directs pour brouillons, révision, publication et types de contenu.",
    "Taxonomies partagées entre le menu public et l’administration.",
    "Création reliée à la destination publique puis ouverture immédiate de l’éditeur.",
    "Éditeur avec identité, contenu, distribution, SEO, sources et score de préparation.",
    "Aperçu privé protégé et actions sauvegarder, réviser, publier et archiver.",
    "Historique de révisions et journal d’actions dans PostgreSQL avec RLS administrateur.",
    "Index de filtrage administrateur et en-têtes HTTP de sécurité.",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("15.5 Prochain ordre de réalisation", level=2)
priorities = [
    ("P0 — verticale plante PostgreSQL", [
        "Étendre le schéma plante avec les champs horticoles, relations, sources et média principal.",
        "Construire l’éditeur plante typé.",
        "Lire la fiche publique depuis PostgreSQL et revalider après publication.",
        "Supprimer la duplication avec les données statiques de lib/plants.ts.",
        "Tester RLS, actions serveur, rendu public et parcours navigateur.",
    ]),
    ("P1 — médiathèque complète", [
        "Upload sécurisé vers le bucket brouillon avec validation MIME, poids et dimensions.",
        "Alt text, crédit, licence, recadrage et sélection depuis l’éditeur.",
        "Promotion contrôlée vers le bucket public lors de la publication.",
    ]),
    ("P1 — workflow et bibliothèque", [
        "Restauration et comparaison des versions.",
        "Pagination et filtres PostgreSQL côté serveur.",
        "Sélection multiple, duplication, archivage et alertes de qualité.",
        "Publication programmée et calendrier éditorial réel.",
    ]),
    ("P2 — recherche et qualité publique", [
        "Brancher la recherche sur PostgreSQL Full Text Search.",
        "Générer breadcrumbs, canonical et JSON-LD depuis les contenus publiés.",
        "Ajouter tests de navigation, accessibilité, performance et monitoring.",
    ]),
]
for title, items in priorities:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    for text in items:
        doc.add_paragraph(text, style="List Bullet 2" if "List Bullet 2" in [s.name for s in doc.styles] else "List Bullet")

doc.add_heading("15.6 Décision de cycle", level=2)
p = doc.add_paragraph()
run = p.add_run("Prochaine verticale obligatoire. ")
run.bold = True
p.add_run(
    "Créer une fiche plante depuis le studio, ajouter contenu structuré, image et sources, "
    "prévisualiser, soumettre en révision, publier, puis afficher exactement la même donnée "
    "sur la fiche publique. Ne pas multiplier les pages statiques avant cette preuve de bout en bout."
)

doc.add_heading("15.7 Références du benchmark", level=2)
for text in [
    "Next.js — React hydration error: https://nextjs.org/docs/messages/react-hydration-error",
    "Sanity — Studio tools: https://www.sanity.io/docs/studio/studio-tools",
    "Sanity — Structure tool: https://www.sanity.io/docs/studio/structure-introduction",
    "Sanity — Content releases: https://www.sanity.io/docs/studio/content-releases",
    "Directus — Content module: https://docs.directus.io/user-guide/content-module/content",
    "Directus — Content versions: https://docs.directus.io/reference/system/versions",
    "Contentful — Content preview: https://www.contentful.com/developers/docs/tutorials/preview/content-preview/",
    "Supabase — Storage access control: https://supabase.com/docs/guides/storage/security/access-control",
    "Supabase — Full Text Search: https://supabase.com/docs/guides/database/full-text-search",
]:
    doc.add_paragraph(text, style="List Bullet")

doc.core_properties.title = "Natura — Plan détaillé de réalisation v1.1"
doc.core_properties.subject = "Plan directeur actualisé après audit senior et benchmark"
doc.save(output)
print(output)
